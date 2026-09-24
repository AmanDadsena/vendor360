"""The ledger as a workbook, for whoever does the shop's books.

Three sheets — sales, stock, and the credit book — because that is how the
person receiving it will use them: one to add up, one to check against the
shelf, one to chase. Dates are real dates and amounts are real numbers, not
strings that look like them: a workbook whose figures cannot be summed is a
screenshot with extra steps.
"""
from __future__ import annotations

import io
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.db import utcnow
from ..models import Customer, InventoryItem, Transaction, UdhaarEntry, Vendor
from . import udhaar as udhaar_service

HEADER_FILL = PatternFill("solid", fgColor="0B7768")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RUPEES = '"Rs "#,##0'


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for index, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    return ws


def ledger_workbook(db: Session, vendor: Vendor, days: int = 90) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    since = utcnow() - timedelta(days=days)

    # ------------------------------------------------------------- sales
    sales = _sheet(
        wb,
        "Sales",
        ["Date", "Item", "Type", "Quantity", "Unit", "Unit value", "Value", "Source"],
        [12, 24, 10, 10, 8, 12, 14, 10],
    )
    items = {
        i.id: i
        for i in db.scalars(
            select(InventoryItem).where(InventoryItem.vendor_id == vendor.id)
        )
    }
    rows = db.scalars(
        select(Transaction)
        .where(
            Transaction.vendor_id == vendor.id,
            Transaction.occurred_at >= since,
        )
        .order_by(Transaction.occurred_at.desc())
    ).all()
    for txn in rows:
        item = items.get(txn.item_id)
        value = txn.qty * (txn.unit_value or 0)
        sales.append([
            txn.occurred_at.date() if txn.occurred_at else None,
            item.sku_name if item else "Unknown",
            txn.type,
            txn.qty,
            item.unit if item else "",
            txn.unit_value or 0,
            value,
            txn.source,
        ])
    for row in sales.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = RUPEES

    # ------------------------------------------------------------- stock
    stock = _sheet(
        wb,
        "Stock",
        ["Item", "Category", "On hand", "Unit", "Reorder at", "Cost", "Price", "Expires"],
        [24, 14, 10, 8, 11, 10, 10, 12],
    )
    for item in sorted(items.values(), key=lambda i: i.sku_name):
        stock.append([
            item.sku_name,
            item.category,
            item.current_qty,
            item.unit,
            item.reorder_point,
            item.unit_cost or 0,
            item.unit_price or 0,
            item.expires_on.date() if item.expires_on else None,
        ])
    for row in stock.iter_rows(min_row=2, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = RUPEES

    # ------------------------------------------------------------ udhaar
    book = _sheet(
        wb,
        "Udhaar",
        ["Customer", "Phone", "Owed", "Days", "Last entry"],
        [22, 14, 12, 8, 12],
    )
    for balance in udhaar_service.balances(db, vendor):
        last = db.scalar(
            select(UdhaarEntry)
            .where(UdhaarEntry.customer_id == balance.customer.id)
            .order_by(UdhaarEntry.occurred_at.desc())
            .limit(1)
        )
        book.append([
            balance.customer.name,
            balance.customer.phone or "",
            balance.owed,
            balance.days_outstanding,
            last.occurred_at.date() if last and last.occurred_at else None,
        ])
    for row in book.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = RUPEES

    # Named so the person who receives it knows whose books these are.
    wb.properties.title = f"{vendor.store_name} — ledger"
    wb.properties.creator = "Vendor360"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


__all__ = ["ledger_workbook", "Customer"]
